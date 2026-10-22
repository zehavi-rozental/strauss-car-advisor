"""When the transport representative is asked, and what the employee is told (stage 2 of the root-cause plan, 22/10).

The rules (A–D): data the employee gives is asked from the employee, never the representative (A); values the rules
compute (budget, use value, electric refund, maintenance by the occupations table, the family-track split) are never
sent for approval (B). The representative is asked only: the copay of every new car compared, calculated meanwhile with
the catalog default (C1); a car with no catalog copay, then not calculated (C2); a role not in the occupations table /
an unclear opt-out (C3); a stated group that contradicts the occupations table (C4); a concrete arrangement question no
source answers (C5). Only code opens a question, never twice for the same item; an answer changes only its own option,
survives a recalculation and a reload, and a future effective date is never shown as in force today (D).
"""
from __future__ import annotations

import shutil
from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from app.modules.cases.public import MissingField, Mode
from app.modules.conversation.llm.fake import FakeLLM, call
from app.modules.conversation.public import ConversationService, LLMChain
from app.modules.conversation.stage import Stage
from app.modules.conversation.templates import UNANSWERED
from app.modules.rules.public import RulesStore

from .conftest import CONFIG, make_client

TOKEN = "rep-policy"
AUTH = {"X-Rep-Token": TOKEN}
PROFILE_B = ["קבוצה ב, שכר בסיס 18,000",
             "אין עמלות, אין בונוס, פנסיה 6%, קרן השתלמות 2.5%, 2.75 נקודות זיכוי, 2 ילדים"]
NO_CAR = ["אין רכב", "אין"]
THREE = ["CAR-05", "CAR-02", "CAR-01"]            # NIRO 250, YARIS CROSS HSD 200, PICANTO 150 (catalog)
CATALOG = {"CAR-05": "250", "CAR-02": "200", "CAR-01": "150", "CAR-10": "300"}


class Chat:
    def __init__(self, sessions, llm: FakeLLM | None = None, config=CONFIG):
        self.llm = llm
        self.svc = ConversationService(sessions, RulesStore(config), LLMChain([llm]) if llm else LLMChain())
        turn, self.token = self.svc.start()
        self.cid = turn.conversation_id

    def say(self, text, *calls, reply=None):
        if self.llm is not None and (calls or reply is not None):
            self.llm.turn(*calls, reply=reply)
        return self.svc.send(self.cid, self.token, text)

    def choose(self, ids, field="options"):
        return self.svc.send_choice(self.cid, self.token, ids, field)

    @property
    def state(self):
        return self.svc.view(self.cid, self.token).state

    def open(self):
        return [q for q in self.svc.cases.list_open() if q.conversation_id == self.cid]


class Http:
    """The employee and the representative through the HTTP API (template mode)."""

    def __init__(self, engine):
        self.c = make_client(engine, rep_access_token=TOKEN)
        self.cid = self.c.post("/api/conversations").json()["conversation_id"]

    def say(self, text):
        return self.c.post(f"/api/conversations/{self.cid}/messages", json={"text": text}).json()

    def choose(self, ids):
        return self.c.post(f"/api/conversations/{self.cid}/messages", json={"choices": ids, "field": "options"}).json()

    def view(self):
        return self.c.get(f"/api/conversations/{self.cid}").json()

    def questions(self):
        return self.c.get("/api/rep/questions", headers=AUTH).json()

    def answer(self, qid, value, **extra):
        return self.c.post(f"/api/rep/questions/{qid}/answer", json={"value": value, **extra}, headers=AUTH)

    def last_bot(self):
        msgs = self.c.get(f"/api/conversations/{self.cid}/updates").json()["messages"]
        return [m for m in msgs if m["role"] == "assistant"][-1]["content"]


def results(view):
    return {r["option_id"]: r["result"] for r in view["results"]}


def group_b_with_three(h: Http):
    for m in [*PROFILE_B, *NO_CAR]:
        h.say(m)
    return h.choose(THREE)


# ---------- C1: every new car compared: calculated with the catalog copay, the representative confirms it ----------

def test_c1_three_cars_three_questions_with_the_catalog_amount_and_a_clear_explanation(engine):
    h = Http(engine)
    turn = group_b_with_three(h)
    qs = h.questions()
    assert sorted(q["option_id"] for q in qs) == sorted(THREE)
    for q in qs:
        assert (q["missing_field"], q["mode"], q["suggested_value"]) == ("NEW_CAR_COPAY", "CONFIRM", CATALOG[q["option_id"]])
        assert CATALOG[q["option_id"]] in q["question_text"] and q["context"]["car_label"] in q["question_text"]
        assert set(q["context"]) <= {"group", "car_label", "catalog_code", "suggested_amount"}   # B and privacy
    res = results(h.view())
    for oid in THREE:
        assert res[oid]["copay_source"] == "catalog" and res[oid]["copay_pending"]
        assert Decimal(str(res[oid]["alternative"]["copay"])) == Decimal(CATALOG[oid])
    reply = turn["reply"]
    for oid, amount in CATALOG.items():                              # every amount used, said in the chat
        if oid in THREE:
            assert f"{amount} ₪" in reply
    assert "קטלוג" in reply and "נציג התחבורה" in reply
    assert "סכום ההשתתפות שלכם ברכב החדש" not in reply             # never as if the employee gave it


def test_c1_the_challenge_script_rep_confirms_350_from_a_future_date_and_the_comparison_updates(engine):
    h = Http(engine)
    for m in ["קבוצה ג, שכר בסיס 14,200", "עמלות 1,000, אין בונוס, פנסיה 6%, קרן השתלמות 2.5%, 2.25 נקודות זיכוי, אין ילדים",
              "יש לי קורולה קרוס, משלמת 300", "אטו 3"]:
        h.say(m)
    before = results(h.view())["CAR-10"]
    assert before["copay_source"] == "catalog"
    [q] = h.questions()
    jan_1 = date(date.today().year + 1, 1, 1)                        # the script's "1/1": always ahead of today
    assert h.answer(q["id"], 350, effective_date=jan_1.isoformat()).status_code == 200
    after = results(h.view())["CAR-10"]                              # a fresh GET: the reload
    assert after["copay_source"] == "representative" and not after["copay_pending"]
    assert after["copay_effective_date"] == jan_1.isoformat()
    assert Decimal(str(after["alternative"]["copay"])) == Decimal("350")
    assert after["delta_display"] != before["delta_display"]
    text = h.last_bot()
    assert "350 ₪" in text and f"ההשתתפות תחול מ-{jan_1:%d.%m.%Y}" in text   # future: never "in force" today
    assert "בתוקף" not in text
    assert after["alternative"]["net_display"] in text.replace(",", "")
    assert h.answer(q["id"], 360).status_code == 409                  # answered once


def test_c1_a_different_amount_changes_only_its_own_car_and_survives_a_recalculation(engine):
    h = Http(engine)
    group_b_with_three(h)
    first = results(h.view())
    picanto = next(q for q in h.questions() if q["option_id"] == "CAR-01")
    assert h.answer(picanto["id"], 180, effective_date="2026-01-01").status_code == 200
    second = results(h.view())
    assert second["CAR-01"]["copay_source"] == "representative"
    assert "ההשתתפות בתוקף מ-01.01.2026" in h.last_bot()             # a past date: in force
    for oid in ("CAR-05", "CAR-02"):                                  # the others are untouched
        assert second[oid] == first[oid]
    h.say("טעיתי, השכר שלי 25,000")                                   # a recalculation of every option
    third = results(h.view())
    assert third["CAR-01"]["today"]["net_display"] != second["CAR-01"]["today"]["net_display"]   # recalculated
    assert third["CAR-01"]["copay_source"] == "representative"           # the approval is kept (D6)
    assert Decimal(str(third["CAR-01"]["alternative"]["copay"])) == Decimal("180")


def test_c1_comparing_all_options_keeps_the_choice_and_opens_no_extra_question(engine):
    h = Http(engine)
    group_b_with_three(h)
    turn = h.say("השווה את כל האפשרויות שלי")
    view = h.view()
    assert view["selected"] == THREE                                  # the employee's choice is kept
    assert sorted(q["option_id"] for q in h.questions()) == sorted(THREE)   # no question for the 5 others
    assert "עד לאישור נציג התחבורה" not in turn["reply"]             # never claims a question that was not sent


def test_c1_the_model_cannot_replace_the_choice_by_ranking_on_a_map_request(sessions):
    chat = Chat(sessions, FakeLLM())
    for m in [*PROFILE_B, *NO_CAR]:
        chat.say(m)
    chat.choose(THREE)
    chat.say("השווה את כל האפשרויות שלי", call("rank_options"))
    assert chat.state.selected == THREE
    assert sorted(q.option_id for q in chat.open()) == sorted(THREE)


def test_c1_no_duplicate_questions_when_the_same_cars_are_chosen_again(sessions):
    chat = Chat(sessions)
    for m in [*PROFILE_B, *NO_CAR]:
        chat.say(m)
    chat.choose(THREE)
    chat.choose(THREE)
    chat.say("טעיתי, פנסיה 7%")
    assert sorted(q.option_id for q in chat.open()) == sorted(THREE)


def test_c1_a_copay_the_employee_expects_for_a_new_car_is_kept_apart_and_sent_for_checking(sessions):
    chat = Chat(sessions)
    for m in [*PROFILE_B, *NO_CAR]:
        chat.say(m)
    chat.choose(["CAR-01"])
    turn = chat.say("בפיקנטו אני אשלם 300")
    s = chat.state
    assert s.current.copay is None                                    # never the copay of a car today
    r = s.results[0].result
    assert r["copay_source"] == "catalog" and Decimal(str(r["alternative"]["copay"])) == Decimal("150")
    assert r["employee_stated_copay"] == "300"                       # shown next to it, never used
    assert "300" in turn.reply and "150" in turn.reply                # the difference is said
    [q] = chat.open()
    assert q.option_id == "CAR-01" and q.context.get("employee_stated_amount") == "300"


# ---------- C2: a car with no catalog copay: not calculated, the representative is asked ----------

def _no_catalog_copay(tmp_path, code="CAR-01"):
    config = tmp_path / "config"
    shutil.copytree(CONFIG, config)
    wb = openpyxl.load_workbook(config / "strauss-data.xlsx")
    ws = wb["קטלוג רכבים"]
    header = [c.value for c in ws[1]]
    col = header.index("השתתפות עובד ברירת מחדל (דמה)") + 1
    for row in ws.iter_rows(min_row=2):
        if row[0].value == code:
            row[col - 1].value = None
    wb.save(config / "strauss-data.xlsx")
    return config


@pytest.mark.parametrize("mode", ["template", "llm"])
def test_c2_a_car_with_no_catalog_copay_is_not_calculated_and_the_rep_is_asked(sessions, tmp_path, mode):
    chat = Chat(sessions, FakeLLM() if mode == "llm" else None, config=_no_catalog_copay(tmp_path))
    for m in [*PROFILE_B, *NO_CAR]:
        chat.say(m)
    turn = chat.choose(["CAR-01", "CAR-05"])
    s = chat.state
    by = {r.option_id: r.result for r in s.results}
    assert by["CAR-05"]["status"] == "OK"                             # the other car is calculated
    assert by["CAR-01"]["status"] != "OK" and by["CAR-01"].get("delta_display") is None   # never invented
    assert "חסר לי סכום ההשתתפות" in turn.reply and "PICANTO" in turn.reply
    assert s.stage == Stage.PRESENTED                                 # the conversation goes on
    q = next(q for q in chat.open() if q.option_id == "CAR-01")
    assert q.missing_field == MissingField.NEW_CAR_COPAY and q.suggested_value is None


def test_c2_under_the_catalog_default_policy_a_missing_amount_is_still_asked_never_only_said(sessions, tmp_path):
    """D8: "שאלתי את נציג התחבורה" only with a real question, also when the catalog amount is otherwise final."""
    config = _no_catalog_copay(tmp_path)
    ext = config / "extensions.yaml"
    ext.write_text(ext.read_text(encoding="utf-8").replace(
        "copay_policy: CATALOG_DEFAULT_CONFIRM_ASYNC", "copay_policy: CATALOG_DEFAULT"), encoding="utf-8")
    chat = Chat(sessions, config=config)
    for m in [*PROFILE_B, *NO_CAR]:
        chat.say(m)
    turn = chat.choose(["CAR-01", "CAR-05"])
    assert "שאלתי את נציג התחבורה" in turn.reply
    assert [q.option_id for q in chat.open()] == ["CAR-01"]           # the catalog amount of CAR-05 is final


# ---------- C3: a role not in the occupations table and an opt-out: the representative decides ----------

def test_c3_a_role_not_in_the_table_asks_the_rep_whether_opting_out_is_allowed(sessions):
    chat = Chat(sessions, FakeLLM())
    chat.say("אני מעצבת גרפית", call("save_profile", occupation="מעצבת גרפית"))
    chat.say("כן")                                                    # yes, that is my role
    for m in ["קבוצה ב, שכר בסיס 18,000", PROFILE_B[1], "יש לי נירו, משלמת 250", "ויתור על הרכב"]:
        chat.say(m)
    [q] = [q for q in chat.open() if q.missing_field == MissingField.OCCUPATION_OPT_OUT]
    assert q.context.get("occupation_text") == "מעצבת גרפית"


@pytest.mark.parametrize("mode", ["template", "llm"])
def test_c3_a_role_in_the_table_never_asks_the_rep(sessions, mode):
    chat = Chat(sessions, FakeLLM() if mode == "llm" else None)
    for m in ["אני מפתחת תוכנה, שכר בסיס 18,000", PROFILE_B[1], "יש לי נירו, משלמת 250", "ויתור על הרכב"]:
        chat.say(m)
    assert chat.open() == []                                          # the table says: allowed, 1,500 (B)
    assert chat.state.results[0].result["status"] == "OK"


# ---------- C4: a stated group that contradicts the occupations table ----------

@pytest.mark.parametrize("messages", [["קבוצה ג, שכר בסיס 18,000", "אני מפתחת תוכנה"],
                                      ["אני מפתחת תוכנה בקבוצה ג, שכר בסיס 18,000"]])
def test_c4_a_group_that_contradicts_the_role_is_never_decided_silently(engine, messages):
    h = Http(engine)
    replies = [h.say(m)["reply"] for m in messages]
    assert any("ג" in r and "ב" in r and "נציג" in r for r in replies)   # both groups named, the rep is asked
    for m in [PROFILE_B[1], "יש לי נירו, משלמת 250"]:
        h.say(m)
    qs = [q for q in h.questions() if q["missing_field"] == "ELIGIBILITY_GROUP"]
    assert len(qs) == 1
    q = qs[0]
    assert q["context"]["stated_group"] == "ג" and q["context"]["table_group"] == "ב"
    assert "18" not in str(q["context"]) and "18,000" not in q["question_text"]   # no pay data
    assert h.view()["stage"] == "WAITING_FOR_REP"                      # no calculation on a guessed group
    assert h.answer(q["id"], "ב").status_code == 200
    view = h.view()
    assert view["state"]["profile"]["group"] == "ב" if "state" in view else True
    assert view["stage"] in ("CHOOSE_OPTIONS", "PRESENTED")


def test_c4_an_answer_that_is_not_a_group_is_refused(engine):
    h = Http(engine)
    for m in ["קבוצה ג, שכר בסיס 18,000", "אני מפתחת תוכנה", PROFILE_B[1], "יש לי נירו, משלמת 250"]:
        h.say(m)
    q = next(q for q in h.questions() if q["missing_field"] == "ELIGIBILITY_GROUP")
    assert h.answer(q["id"], "קבוצה כלשהי").status_code == 422


def test_c4_the_model_cannot_decide_the_group_either(sessions):
    chat = Chat(sessions, FakeLLM())
    chat.say("קבוצה ג, שכר בסיס 18,000", call("save_profile", eligibility_group="ג", base_salary=18000))
    chat.say("אני מפתחת תוכנה", call("save_profile", occupation="מפתחת תוכנה"))
    s = chat.state
    assert s.group_conflict is not None and s.group_conflict.stated == "ג" and s.group_conflict.table == "ב"


# ---------- C5: process and result questions are answered here, never sent ----------

@pytest.mark.parametrize("question", ["למה צריך את הנתון הזה?", "חייבים לכתוב עמלות?",
                                      "שאלתי האם חיבים לכתוב שיש עמלות", "למה אתה שואל על ילדים?"])
@pytest.mark.parametrize("mode", ["template", "llm"])
def test_c5_a_question_about_the_employees_own_data_is_answered_not_sent(sessions, question, mode):
    chat = Chat(sessions, FakeLLM() if mode == "llm" else None)
    chat.say(PROFILE_B[0])
    turn = chat.say(question)
    assert UNANSWERED not in turn.reply and chat.state.unanswered_question is None
    assert chat.open() == []


def test_a_why_question_while_todays_copay_is_unknown_reads_well(sessions):
    """Rule A wording: "החישוב מחכה רק להשתתפות…" (ל + ה), never "לההשתתפות"."""
    chat = Chat(sessions)
    for m in [PROFILE_B[0], *PROFILE_B[1:], "יש לי קורולה קרוס", "לא יודעת", "PICANTO"]:
        chat.say(m)
    turn = chat.say("למה צריך את הנתון הזה?")
    assert "לההשתתפות" not in turn.reply and chat.open() == []
    assert "החישוב מחכה רק להשתתפות ברכב של היום" in turn.reply


@pytest.mark.parametrize("question", ["למה ההערכה זמנית?", "למה החישוב זמני?", "זה סופי?", "מה זה אומר הערכה?",
                                      "הנציג כבר קיבל את השאלה?"])
@pytest.mark.parametrize("mode", ["template", "llm"])
def test_c5_a_question_about_the_result_is_answered_from_it(sessions, question, mode):
    chat = Chat(sessions, FakeLLM() if mode == "llm" else None)
    for m in [*PROFILE_B, *NO_CAR]:
        chat.say(m)
    chat.choose(["CAR-01"])
    turn = chat.say(question)
    assert UNANSWERED not in turn.reply
    assert "150 ₪" in turn.reply and "PICANTO" in turn.reply          # the amount used and for which car
    assert "ברירת המחדל בקטלוג" in turn.reply and "מאשר את ההשתתפות שתחול עליכם בפועל" in turn.reply
    assert "אילו רכבים לבדוק" not in turn.reply                      # never back to choosing cars
    assert [q.option_id for q in chat.open()] == ["CAR-01"]           # nothing new sent


def test_c5_a_concrete_arrangement_question_is_still_offered_to_the_rep(sessions):
    chat = Chat(sessions)
    chat.say(PROFILE_B[0])
    turn = chat.say("האם מותר לנסוע עם רכב החברה לחו\"ל?")
    assert UNANSWERED in turn.reply


# ---------- A: data the employee gives is asked from the employee, never the representative ----------

@pytest.mark.parametrize("mode", ["template", "llm"])
def test_a_an_unknown_copay_of_the_car_today_is_asked_from_the_employee(sessions, mode):
    chat = Chat(sessions, FakeLLM() if mode == "llm" else None)
    for m in [*PROFILE_B, "יש לי נירו"]:
        chat.say(m)
    turn = chat.say("לא ידוע", *([call("save_profile", current_copay_unknown=True)] if mode == "llm" else []))
    assert chat.open() == []
    assert chat.state.stage != Stage.WAITING_FOR_REP
    chat.choose(["CAR-01"])
    turn = chat.say("מה עכשיו?") if mode == "template" else turn
    assert chat.open() == [] and not chat.state.results or chat.state.results[0].result.get("status") != "OK"
    turn = chat.say("250")
    s = chat.state
    assert s.current.copay == Decimal("250") and s.results and s.results[0].result["status"] == "OK"


def test_a_the_rep_never_sees_a_question_about_the_employees_pay(engine):
    h = Http(engine)
    group_b_with_three(h)
    for q in h.questions():
        blob = q["question_text"] + str(q["context"])
        assert "18,000" not in blob and "18000" not in blob and "פנסיה" not in blob


# ---------- B: values the rules compute are never sent for approval ----------

def test_b_rule_values_never_go_to_the_rep(sessions):
    chat = Chat(sessions)
    for m in ["אני מנהלת מוצר, שכר בסיס 18,500, אין עמלות ואין בונוס, פנסיה 6%, קרן השתלמות 2.5%, 2.75 נקודות זיכוי, 5 ילדים",
              "יש לי קורולה קרוס, משלמת 300"]:
        chat.say(m)
    chat.choose(["CAR-10", "CAR-19"])                                 # electric (refund) + family track (split)
    fields = {q.missing_field for q in chat.open()}
    assert fields == {MissingField.NEW_CAR_COPAY}
    for q in chat.open():
        assert set(q.context) <= {"group", "car_label", "catalog_code", "suggested_amount"}


# ---------- D: only code opens a question; privacy and permissions ----------

def test_d_the_model_cannot_open_a_rep_question_by_itself(sessions):
    chat = Chat(sessions, FakeLLM())
    chat.say(PROFILE_B[0])
    chat.say("למה אתה צריך את כל זה?", call("propose_rep_question", question="למה צריך את כל הנתונים"))
    assert chat.open() == [] and chat.state.unanswered_question is None
    chat.say("האם מותר לנסוע עם רכב החברה לחו\"ל?", call("propose_rep_question", question="מותר לנסוע לחו\"ל?"))
    assert chat.open() == []                                          # an offer at most; nothing sent without "לשלוח"


def test_d_the_rep_api_needs_its_token(engine):
    h = Http(engine)
    group_b_with_three(h)
    assert h.c.get("/api/rep/questions").status_code == 401
    q = h.questions()[0]
    assert h.c.post(f"/api/rep/questions/{q['id']}/answer", json={"value": 1}).status_code == 401
    assert "conversation_id" not in q


def test_d_an_answer_is_applied_once_and_kept_after_reload(engine):
    h = Http(engine)
    group_b_with_three(h)
    q = next(q for q in h.questions() if q["option_id"] == "CAR-02")
    h.answer(q["id"], 200)
    one, two = results(h.view()), results(h.view())
    assert one == two and one["CAR-02"]["copay_source"] == "representative"
    assert one["CAR-02"]["copay_effective_date"] is None
    assert date.today().isoformat()                                   # (the reload is the second GET)


# ---------- a car taken out of the comparison: its open question is cancelled, an answered one is kept ----------

def test_a_car_taken_out_of_the_comparison_cancels_only_its_open_question(sessions):
    chat = Chat(sessions)
    for m in [*PROFILE_B, *NO_CAR]:
        chat.say(m)
    chat.choose(["CAR-01", "CAR-05"])
    by = {q.option_id: q for q in chat.open()}
    chat.svc.cases.answer(by["CAR-05"].id, "260")                     # NIRO answered
    chat.choose(["CAR-02"])                                           # both taken out
    assert [q.option_id for q in chat.open()] == ["CAR-02"]
    assert chat.svc.cases.get(by["CAR-01"].id).status.value == "CANCELLED"   # kept as history, not deleted
    assert chat.svc.cases.get(by["CAR-05"].id).status.value == "ANSWERED"    # never touched
    chat.choose(["CAR-05"])                                           # chosen again: the answer is used
    r = chat.state.results[0].result
    assert r["copay_source"] == "representative" and r["alternative"]["copay"] == "260"
    assert chat.open() == []
