"""טעינת קבצי ההגדרות: strauss-rules.json, extensions.yaml, strauss-data.xlsx.

מצב strict: כל שגיאה עוצרת את העלייה.
מצב dev: קובץ שגוי בטעינה חוזרת → נשארים עם הגרסה הקודמת ומדווחים אזהרה.
"""
from __future__ import annotations

import hashlib
import json
import logging
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Literal

import yaml
from openpyxl import load_workbook
from pydantic import ValidationError

from app.modules.calculation.public import FAMILY_GROUP, Bracket, CalcRules, TwoTier

from .faq import Faq, faq_values, load_faq
from .models import (
    DRIVE_FROM_HEBREW,
    MOVED_TO_RULES,
    CatalogCar,
    Extensions,
    Occupation,
    ReferenceEmployee,
    ReferenceResult,
    StraussRules,
)

log = logging.getLogger(__name__)

RULES_FILE = "strauss-rules.json"
EXTENSIONS_FILE = "extensions.yaml"
DATA_FILE = "strauss-data.xlsx"

MAX_COMPARE = 3          # alternatives compared at once (conversation.stage.MAX_OPTIONS)

SHEET_CATALOG = "קטלוג רכבים"
SHEET_OCCUPATIONS = "עיסוקים"
SHEET_EMPLOYEES = "עובדים לדוגמה"
SHEET_RESULTS = "חישובים"


class ConfigError(Exception):
    pass


# בסיס ההפרשות לפנסיה ולקרן השתלמות (netMethod.contributionsBase)
CONTRIBUTION_COMPONENTS = {
    "gross": ["base", "commissions", "bonus"],
    "base": ["base"],
}


def _schema_error(file_name: str, exc: ValidationError) -> ConfigError:
    """שגיאת סכמה בעברית: שם הקובץ ונתיב השדה (כמו בקובץ), בלי מונחים של pydantic."""
    problems = []
    for err in exc.errors():
        path = ".".join(str(part) for part in err["loc"])
        if err["type"] == "missing":
            problems.append(f"חסר השדה {path}")
        elif err["type"] in ("literal_error", "enum"):
            expected = err.get("ctx", {}).get("expected", "")
            problems.append(f"ערך לא נתמך ב-{path}: {err['input']!r} (אפשרי: {expected})")
        else:
            problems.append(f"ערך לא תקין ב-{path}: {err['msg']}")
    return ConfigError(f"{file_name}: " + "; ".join(problems))


def _check_family_cars(strauss: StraussRules, catalog: list[CatalogCar]) -> None:
    """familyTrack.cars הוא המסלול המשפחתי — בדיוק: כל רכב ברשימה בקטלוג ובקבוצה משפחתי, ולהפך."""
    by_code = {c.catalog_code: c for c in catalog}
    listed = strauss.family_track.cars
    for code in listed:
        if code not in by_code:
            raise ConfigError(f"הרכב {code} ברשימת המסלול המשפחתי (familyTrack.cars) אינו בקטלוג")
        if by_code[code].group != FAMILY_GROUP:
            raise ConfigError(f"הרכב {code} ברשימת המסלול המשפחתי אינו בקבוצה '{FAMILY_GROUP}' בקטלוג")
    extra = sorted(c.catalog_code for c in catalog if c.group == FAMILY_GROUP and c.catalog_code not in listed)
    if extra:
        raise ConfigError(f"רכבים בקבוצה '{FAMILY_GROUP}' בקטלוג שאינם ב-familyTrack.cars: {extra}")


def _parse_extensions(raw: Any) -> Extensions:
    moved = [k for k in MOVED_TO_RULES if isinstance(raw, dict) and k in raw]
    if moved:
        where = ", ".join(f"{k} → {MOVED_TO_RULES[k]}" for k in moved)
        raise ConfigError(f"{EXTENSIONS_FILE}: השדות עברו לקובץ הכללים {RULES_FILE} — יש להסיר אותם: {where}")
    try:
        return Extensions.model_validate(raw)
    except ValidationError as exc:
        raise _schema_error(EXTENSIONS_FILE, exc) from exc


@dataclass(frozen=True)
class RulesSnapshot:
    strauss: StraussRules
    extensions: Extensions
    calc: CalcRules
    catalog: list[CatalogCar]
    occupations: dict[str, Occupation]
    reference_employees: list[ReferenceEmployee]
    reference_results: list[ReferenceResult]
    file_hashes: dict[str, str]
    file_contents: dict[str, bytes] = field(repr=False, default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    faq: Faq = field(default_factory=Faq)          # faq.yaml (answers only; not part of the calculation)


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()[:16]


def _money(v: Any) -> Decimal:
    if v is None or v == "":
        raise ConfigError("ערך כספי חסר")
    return Decimal(str(v))


def _yes(v: Any) -> bool:
    return str(v).strip() in ("כן", "true", "True", "1")


def _rows(ws) -> Iterable[dict[str, Any]]:
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    for row in it:
        if row is None or all(c is None for c in row):
            continue
        yield dict(zip(header, row))


def build_calc_rules(strauss: StraussRules, ext: Extensions) -> CalcRules:
    nm = strauss.net_method

    def tier(t) -> TwoTier:
        return TwoTier(threshold=t.threshold, ceiling=t.ceiling, low_rate=t.low_rate, high_rate=t.high_rate)

    return CalcRules(
        rules_version=strauss.version,
        extensions_version=ext.version,
        budgets=strauss.budgets,
        group_order=strauss.group_order,
        electric_net_bonus=strauss.electric.net_bonus,
        electric_method=strauss.electric.method,
        family_min_children=strauss.family_track.min_children,
        family_gap_employee_share=strauss.family_track.gap_split_employee,
        family_gap_adds_to_copay=ext.family_gap_adds_to_copay,
        plug_in_counts_as_electric=ext.plug_in_counts_as_electric,
        copay_reduces_use_value=nm.participation.reduces_value_of_use,
        tax_year=nm.tax_year,
        credit_point_value=nm.credit_point_value,
        brackets=[Bracket(up_to=u, rate=r) for u, r in nm.brackets],
        national_insurance=tier(nm.national_insurance),
        health=tier(nm.health),
        pension_credit_rate=nm.pension_credit.rate,
        pension_credit_income_ceiling=nm.pension_credit.income_ceiling,
        pension_credit_max_rate=nm.pension_credit.max_contribution_rate,
        pensionable_components=CONTRIBUTION_COMPONENTS[nm.contributions_base],
        study_fund_components=CONTRIBUTION_COMPONENTS[nm.contributions_base],
    )


def _load_workbook(path: Path, ext: Extensions, strauss: StraussRules):
    wb = load_workbook(path, data_only=True, read_only=True)
    for sheet in (SHEET_CATALOG, SHEET_OCCUPATIONS, SHEET_EMPLOYEES, SHEET_RESULTS):
        if sheet not in wb.sheetnames:
            raise ConfigError(f"חסר גיליון '{sheet}' בקובץ {path.name}")

    catalog: list[CatalogCar] = []
    for r in _rows(wb[SHEET_CATALOG]):
        drive_text = str(r["הנעה"]).strip()
        if drive_text not in DRIVE_FROM_HEBREW:
            raise ConfigError(f"הנעה לא מוכרת '{drive_text}' ברכב {r['מזהה']}")
        catalog.append(CatalogCar(
            catalog_code=str(r["מזהה"]).strip(),
            group=str(r["קבוצה"]).strip(),
            manufacturer=str(r["תוצר"]).strip(),
            manufacturer_code=str(r["קוד תוצר"]).strip().zfill(5),
            model_name=str(r["דגם (כינוי מסחרי)"]).strip(),
            model_code=str(r["קוד דגם"]).strip().zfill(5),
            model_description=str(r["תאור דגם"]).strip(),
            reg_year=int(r["שנת רישום"]),
            vehicle_type=ext.default_vehicle_type,
            drive=DRIVE_FROM_HEBREW[drive_text],
            list_price=_money(r["מחיר מחירון"]),
            use_value=_money(r["שווי שימוש חודשי"]),
            monthly_cost=_money(r["עלות חודשית לחברה (דמה)"]),
            default_copay=_money(r["השתתפות עובד ברירת מחדל (דמה)"]),
        ))

    known_groups = set(strauss.group_order) | {"משפחתי"}
    bad = [c.catalog_code for c in catalog if c.group not in known_groups]
    if bad:
        raise ConfigError(f"קבוצה לא מוכרת ברכבים {bad}")
    unknown = sorted(set(ext.car_aliases) - {c.catalog_code for c in catalog})
    if unknown:
        raise ConfigError(f"כינויים לרכבים שאינם בקטלוג: {unknown}")
    catalog = [c.model_copy(update={"aliases": tuple(ext.car_aliases.get(c.catalog_code, ()))}) for c in catalog]
    _check_family_cars(strauss, catalog)

    occupations: dict[str, Occupation] = {}
    for r in _rows(wb[SHEET_OCCUPATIONS]):
        occ = Occupation(
            name=str(r["עיסוק"]).strip(),
            group=str(r["קבוצת זכאות"]).strip(),
            car_mandatory=_yes(r["רכב חובה"]),
            opt_out_allowed=_yes(r["מותר לוותר"]),
            maintenance=_money(r["אחזקת רכב בוויתור (₪)"] or 0),
        )
        if occ.opt_out_allowed and occ.maintenance <= 0:
            raise ConfigError(f"לעיסוק '{occ.name}' מותר ויתור אבל אין סכום אחזקה")
        occupations[occ.name] = occ

    employees = [ReferenceEmployee(
        employee_id=str(r["מזהה"]).strip(), name=str(r["שם"]).strip(),
        occupation=str(r["עיסוק"]).strip(), group=str(r["קבוצה"]).strip(),
        base=_money(r["שכר בסיס"]), commissions=_money(r["עמלות"] or 0), bonus=_money(r["בונוס"] or 0),
        pension_percent=_money(r["פנסיה עובד %"]), study_fund_percent=_money(r["קרן השתלמות עובד %"] or 0),
        credit_points=_money(r["נקודות זיכוי"]), children=int(r["ילדים עד 23"] or 0),
        current_car=str(r["רכב היום"]).strip(), current_copay=_money(r["השתתפות היום"] or 0),
        question=str(r["השאלה"]).strip(),
    ) for r in _rows(wb[SHEET_EMPLOYEES])]   # עמודת ת"ז לא נקראת בכוונה

    results = [ReferenceResult(
        employee_id=str(r["מזהה"]).strip(), today_label=str(r["היום"]).strip(),
        alternative_label=str(r["החלופה"]).strip(), eligible=_yes(r["זכאות"]), reason=str(r["סיבה"]).strip(),
        net_today=_money(r["נטו היום"]), net_alternative=_money(r["נטו בחלופה"]), delta=_money(r["הפרש"]),
        tax_today=_money(r["מס היום"]), tax_alternative=_money(r["מס בחלופה"]),
        social_today=_money(r["ביטוח לאומי+בריאות היום"]),
        social_alternative=_money(r["ביטוח לאומי+בריאות בחלופה"]),
        electric_gross=_money(r["החזר חשמל ברוטו"] or 0),
        family_share=_money(r["חלק העובד במסלול המשפחתי"] or 0),
    ) for r in _rows(wb[SHEET_RESULTS])]
    wb.close()
    return catalog, occupations, employees, results


def load_snapshot(config_dir: Path) -> RulesSnapshot:
    """טוען ומאמת את שלושת הקבצים. זורק ConfigError על כל בעיה."""
    paths = {name: config_dir / name for name in (RULES_FILE, EXTENSIONS_FILE, DATA_FILE)}
    for name, p in paths.items():
        if not p.exists():
            raise ConfigError(f"חסר קובץ הגדרות: {p}")
    contents = {name: p.read_bytes() for name, p in paths.items()}
    try:
        try:
            strauss = StraussRules.model_validate(json.loads(contents[RULES_FILE].decode("utf-8")))
        except ValidationError as exc:
            raise _schema_error(RULES_FILE, exc) from exc
        ext = _parse_extensions(yaml.safe_load(contents[EXTENSIONS_FILE].decode("utf-8")))
        catalog, occupations, employees, results = _load_workbook(paths[DATA_FILE], ext, strauss)
        calc = build_calc_rules(strauss, ext)
        faq = load_faq(config_dir, faq_values(strauss, MAX_COMPARE))
    except ConfigError:
        raise
    except Exception as exc:  # שגיאות סכמה/פענוח → הודעה אחידה
        raise ConfigError(str(exc)) from exc

    return RulesSnapshot(
        strauss=strauss, extensions=ext, calc=calc, catalog=catalog, occupations=occupations,
        reference_employees=employees, reference_results=results,
        file_hashes={n: _sha(b) for n, b in contents.items()}, file_contents=contents, faq=faq,
    )


class RulesStore:
    """מחזיק את הגרסה הפעילה של ההגדרות ויודע לטעון מחדש לפי המצב (dev/strict)."""

    def __init__(self, config_dir: Path, mode: Literal["dev", "strict"] = "dev"):
        self.config_dir = config_dir
        self.mode = mode
        self._current: RulesSnapshot | None = None
        self.last_error: str | None = None

    @property
    def current(self) -> RulesSnapshot:
        if self._current is None:
            self.reload()
        assert self._current is not None
        return self._current

    def reload(self) -> RulesSnapshot:
        try:
            snap = load_snapshot(self.config_dir)
        except ConfigError as exc:
            self.last_error = str(exc)
            if self.mode == "strict" or self._current is None:
                raise
            log.warning("קובץ הגדרות שגוי, נשארים עם הגרסה הקודמת: %s", exc)
            return self._current
        self.last_error = None
        self._current = snap
        return snap
